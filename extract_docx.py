from docx import Document
import os

for i in range(1, 7):
    filepath = f'Bài {i}.docx'
    if os.path.exists(filepath):
        doc = Document(filepath)
        out = f'bai_{i}_content.txt'
        with open(out, 'w', encoding='utf-8') as f:
            f.write(f'===== BÀI {i} =====\n')
            for p in doc.paragraphs:
                if p.text.strip():
                    f.write(p.text + '\n')
            for t_idx, table in enumerate(doc.tables):
                f.write(f'\n--- Table {t_idx + 1} ---\n')
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    f.write(' | '.join(cells) + '\n')
        print(f'Extracted: {out}')
print('Done')
